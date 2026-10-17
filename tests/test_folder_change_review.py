import contextlib
import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import socket
import subprocess

from openpyxl import load_workbook
import folder_change_review as app
from runtime_safety import offline_runtime, safe_operation
from state_store import StateHistory, read_metadata, write_container


class FolderChangeReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name)
        self.root = self.work / "project"
        self.root.mkdir()

    def snapshot(self, name):
        with contextlib.redirect_stdout(io.StringIO()):
            return app.create_snapshot(self.root, self.work / name)

    def test_snapshot_exclusions_archives_and_empty_folders(self):
        for name in ("~note.docx", "~$lock.docx", "THUMBS.DB", "scratch.TMP", "scratch.TeMp"):
            (self.root / name).write_bytes(b"ignored")
        (self.root / "empty").mkdir()
        (self.root / "report.pdf").write_bytes(b"%PDF-1.4\n")
        with zipfile.ZipFile(self.root / "package.zip", "w") as archive:
            archive.writestr("inside.pdf", b"%PDF-1.4\n")
        for name in ("package.rar", "package.7z"):
            (self.root / name).write_bytes(b"ordinary file bytes")
        snapshot = self.snapshot("A.fcrstate")
        self.assertEqual(len(snapshot["files"]), 4)
        self.assertEqual(snapshot["excluded_files"], 5)
        for record in snapshot["files"]:
            path = Path(record["path"])
            self.assertEqual(record["sha256"], hashlib.sha256(path.read_bytes()).hexdigest())
            self.assertEqual(record["parts"], [path.name])
        self.assertEqual(app.load_snapshot(self.work / "A.fcrstate"), snapshot)

    def test_union_hash_changes_moves_and_full_paths(self):
        for directory in ("one", "two"):
            (self.root / directory).mkdir()
            (self.root / directory / "same.pdf").write_bytes(b"same")
        (self.root / "changed.docx").write_bytes(b"before")
        (self.root / "deleted.pdf").write_bytes(b"delete")
        (self.root / "move.pdf").write_bytes(b"move")
        before = self.snapshot("A.fcrstate")
        (self.root / "changed.docx").write_bytes(b"AFTER!")  # same size, different hash
        (self.root / "deleted.pdf").unlink()
        (self.root / "added.pdf").write_bytes(b"add")
        (self.root / "move.pdf").rename(self.root / "one/move.pdf")
        os.utime(self.root / "one/same.pdf", (1_700_000_000, 1_700_000_000))
        after = self.snapshot("B.fcrstate")
        rows = app.compare_snapshots(before, after)
        self.assertEqual(len(rows), 7)
        self.assertEqual({status: sum(row["Status"] == status for row in rows)
                          for status in ("Added", "Deleted", "Changed", "Unchanged")},
                         {"Added": 2, "Deleted": 2, "Changed": 1, "Unchanged": 2})
        for row in rows:
            self.assertEqual(sum(row[name] or 0 for name in ("Added", "Deleted", "Changed")),
                             int(row["Status"] != "Unchanged"))
        self.assertIsNone(next(row for row in rows if row["FileName"] == "added.pdf")["SizeA"])

    def test_excel_plain_numeric_flags_and_literal_names(self):
        (self.root / "=1+1.txt").write_bytes(b"text")
        (self.root / "00123.txt").write_bytes(b"text")
        (self.root / "пример-😀.pdf").write_bytes(b"pdf")
        after = self.snapshot("B.fcrstate")
        rows = app.compare_snapshots({"files": []}, after)
        output = self.work / "changes.xlsx"
        app.write_report(rows, output)
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["Changes"])
        sheet = book["Changes"]
        self.assertEqual(len(sheet.tables), 0)
        headers = [cell.value for cell in sheet[6]][1:]
        self.assertNotIn("Leaf", headers)
        self.assertNotIn("FullPath", headers)
        self.assertFalse(any("SHA" in name.upper() for name in headers))
        self.assertEqual(headers[0], "Row number")
        self.assertEqual(headers, ["Row number", "Level 1", "Leaf depth", "Name without extension",
                                  "File name", "Extension", "Size A (bytes)", "Size A (MiB)", "Size B (bytes)", "Size B (MiB)",
                                  "Size delta (bytes)", "Size delta (MiB)",
                                  "Status", "Added", "Deleted", "Changed", "last"])
        self.assertEqual(headers[-1], "last")
        self.assertEqual(sheet.max_row, 9)
        self.assertEqual(sheet.freeze_panes, "A7")
        self.assertEqual(sheet.sheet_view.zoomScale, 70)
        self.assertEqual(sheet.row_dimensions[6].height, 72)
        self.assertEqual(sheet["B6"].fill.fgColor.rgb, "FF000000")
        self.assertEqual(sheet["C6"].fill.fgColor.rgb, "FF00B0F0")
        self.assertEqual(list(next(sheet.iter_rows(min_row=4, max_row=4, min_col=2, values_only=True))), list(range(1, len(headers) + 1)))
        for cells in sheet.iter_rows(min_row=7, min_col=2):
            self.assertEqual(cells[headers.index("File name")].data_type, "s")
            self.assertEqual(cells[headers.index("Added")].value, 1)
            self.assertEqual(cells[headers.index("Added")].data_type, "n")
            self.assertIsNone(cells[headers.index("Deleted")].value)
            self.assertIsNone(cells[headers.index("Size A (MiB)")].value)
            self.assertEqual(cells[headers.index("Size B (MiB)")].number_format, "#,##0.000")
            self.assertIn(cells[headers.index("Extension")].value, ("txt", "pdf"))
            self.assertEqual(cells[-1].value, 1)
            self.assertTrue(all(cell.fill.patternType is None for cell in cells))
            self.assertTrue(all(cell.border.left.style == "thin" for cell in cells))
            self.assertTrue(all(cell.data_type != "f" and cell.hyperlink is None for cell in cells))

    def test_english_order_keeps_sizes_flags_hierarchy_and_names_aligned(self):
        (self.root / "nested").mkdir()
        (self.root / "nested/данные.tar.gz").write_bytes(b"before")
        (self.root / "deleted.pdf").write_bytes(b"deleted")
        (self.root / ".gitignore").write_bytes(b"")
        before = self.snapshot("A.fcrstate")
        (self.root / "nested/данные.tar.gz").write_bytes(b"x" * 1_234_567)
        (self.root / "deleted.pdf").unlink()
        (self.root / "zero").write_bytes(b"")
        after = self.snapshot("B.fcrstate")
        output = self.work / "aligned.xlsx"
        app.write_report(app.compare_snapshots(before, after), output)
        book = load_workbook(output)
        self.addCleanup(book.close)
        sheet = book["Changes"]
        self.assertEqual(sheet["B2"].value, f"Table: File changes in folder {self.root}")
        headers = [cell.value for cell in sheet[6]][1:]
        self.assertEqual(headers, ["Row number", "Level 1", "Level 2", "Leaf depth", "Name without extension",
                                  "File name", "Extension", "Size A (bytes)", "Size A (MiB)", "Size B (bytes)", "Size B (MiB)",
                                  "Size delta (bytes)", "Size delta (MiB)",
                                  "Status", "Added", "Deleted", "Changed", "last"])
        records = {row[headers.index("File name")]: dict(zip(headers, row))
                   for row in sheet.iter_rows(min_row=7, min_col=2, values_only=True)}
        changed = records["данные.tar.gz"]
        self.assertEqual((changed["Level 1"], changed["Level 2"], changed["Leaf depth"]),
                         ("nested", "данные.tar.gz", 2))
        self.assertEqual((changed["Name without extension"], changed["Extension"]), ("данные.tar", "gz"))
        expected = {"данные.tar.gz": (6, 1_234_567, "Changed", None, None, 1),
                    "deleted.pdf": (7, None, "Deleted", None, 1, None),
                    "zero": (None, 0, "Added", 1, None, None),
                    ".gitignore": (0, 0, "Unchanged", None, None, None)}
        for name, (size_a, size_b, status, added, deleted, changed) in expected.items():
            record = records[name]
            for label, size in (("Size A (MiB)", size_a), ("Size B (MiB)", size_b)):
                if size is None:
                    self.assertIsNone(record[label])
                else:
                    self.assertAlmostEqual(record[label], size / 1_048_576, delta=1e-14)
            self.assertEqual((record["Status"], record["Added"], record["Deleted"], record["Changed"]),
                             (status, added, deleted, changed))
            self.assertEqual(record["last"], 1)
        self.assertEqual(records[".gitignore"]["Name without extension"], ".gitignore")
        self.assertEqual(records[".gitignore"]["Extension"], "")

    def test_failed_snapshot_preserves_previous_output(self):
        (self.root / "file.pdf").write_bytes(b"data")
        output = self.work / "A.fcrstate"
        output.write_bytes(b"previous snapshot")
        with patch.object(app, "write_container", side_effect=PermissionError("unwritable output")):
            with self.assertRaises(PermissionError):
                app.create_snapshot(self.root, output)
        self.assertEqual(output.read_bytes(), b"previous snapshot")

    def test_failed_report_preserves_output_and_cleans_staging(self):
        output = self.work / "changes.xlsx"
        output.write_bytes(b"previous report")
        row = {"parts": ["x" * 32768], "FullPath": str(self.root / ("x" * 32768)), "FileName": "x",
               "Extension": "", "Status": "Added", "Added": 1, "Deleted": 0,
               "Changed": 0, "SizeA": None, "SizeB": 1}
        with self.assertRaises(ValueError):
            app.write_report([row], output)
        self.assertEqual(output.read_bytes(), b"previous report")
        self.assertEqual(list(self.work.glob(".folder-change-review-*")), [])

    def test_no_same_root_check_and_in_project_output_is_allowed(self):
        with contextlib.redirect_stdout(io.StringIO()):
            app.create_snapshot(self.root, self.root / "A.fcrstate")
        self.assertTrue((self.root / "A.fcrstate").exists())
        before = self.snapshot("A.fcrstate")
        other = self.work / "different-project"
        other.mkdir()
        with contextlib.redirect_stdout(io.StringIO()):
            app.create_snapshot(other, self.work / "B.fcrstate")
            counts = app.compare_files(self.work / "A.fcrstate", self.work / "B.fcrstate", self.work / "empty.xlsx")
        self.assertEqual(sum(counts.values()), 0)
        self.assertEqual(before["files"], [])
        book = load_workbook(self.work / "empty.xlsx")
        self.addCleanup(book.close)
        self.assertEqual(book["Changes"].max_row, 6)
        self.assertEqual(book["Changes"]["B2"].value, f"Table: File changes in folder {other}")
        self.assertEqual(book["Changes"]["B6"].value, "Row number")

    def test_invalid_and_duplicate_snapshot_records_are_rejected(self):
        (self.root / "file.txt").write_bytes(b"sample")
        snapshot = self.snapshot("A.fcrstate")
        snapshot["files"].append(dict(snapshot["files"][0]))
        (self.work / "duplicate.json").write_text(json.dumps(snapshot), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Duplicate"):
            app.load_snapshot(self.work / "duplicate.json")

    def test_file_mutation_during_hashing_is_rejected(self):
        path = self.root / "file.txt"
        path.write_bytes(b"sample")
        original = app.os.stat
        def change_before_final_stat(candidate, **kwargs):
            path.write_bytes(b"different length")
            return original(candidate, **kwargs)
        with patch.object(app.os, "stat", side_effect=change_before_final_stat):
            with self.assertRaisesRegex(OSError, "changed while"):
                app.hash_file(str(path))

    def test_lca_uses_union_including_deleted_branch_and_single_file(self):
        paths = [str(self.root / "shared/Finance/a.pdf"), str(self.root / "shared/Financial/b.pdf")]
        rows = [{"FullPath": path, "parts": ["unused"]} for path in paths]
        ancestor, parts = app.report_hierarchy(rows)
        self.assertEqual(ancestor, str(self.root / "shared"))
        self.assertEqual(parts, [["Finance", "a.pdf"], ["Financial", "b.pdf"]])
        self.assertEqual(app.report_hierarchy(rows[:1]), (str(self.root / "shared/Finance"), [["a.pdf"]]))
        windows = [{"FullPath": r"C:\Project\Common\one\a.pdf", "parts": []},
                   {"FullPath": r"C:\Project\Common\two\b.pdf", "parts": []}]
        self.assertEqual(app.report_hierarchy(windows), (r"C:\Project\Common", [["one", "a.pdf"], ["two", "b.pdf"]]))

    def test_parallel_snapshot_matches_serial_file_records(self):
        for index in range(17):
            (self.root / f"file-{index}.pdf").write_bytes(bytes([index]) * (index + 1))
        with contextlib.redirect_stdout(io.StringIO()):
            serial = app.create_snapshot(self.root, self.work / "serial.fcrstate", workers=1)
            for workers in (16, 32):
                parallel = app.create_snapshot(self.root, self.work / "parallel.fcrstate", workers=workers)
                self.assertEqual(serial["files"], parallel["files"])

    def test_machine_state_compact_roundtrip_and_metadata_only_history(self):
        (self.root / "report.pdf").write_bytes(b"document")
        history = StateHistory(self.work)
        output = history.new_state_path()
        with contextlib.redirect_stdout(io.StringIO()):
            state = app.create_snapshot(self.root, output)
        with zipfile.ZipFile(output) as archive:
            self.assertEqual(set(archive.namelist()), {"metadata.json", "files.json", "errors.json"})
            records = json.loads(archive.read("files.json"))
            self.assertEqual(records, [["report.pdf", 8, hashlib.sha256(b"document").hexdigest()]])
        self.assertEqual(app.load_snapshot(output), state)
        self.assertEqual(read_metadata(output)["file_count"], 1)
        # Chooser does not decompress the inventory or need the project online.
        original_read = zipfile.ZipFile.read
        def metadata_only(archive, name, *args, **kwargs):
            self.assertNotEqual(getattr(name, "filename", name), "files.json")
            return original_read(archive, name, *args, **kwargs)
        with patch.object(zipfile.ZipFile, "read", metadata_only):
            items, errors = history.list_states()
        self.assertEqual([item["path"] for item in items], [output])
        self.assertEqual(errors, [])
        history.save_settings({"project": str(self.root), "workers": 32})
        self.assertEqual(history.settings()["workers"], 32)

    def test_corrupt_history_entry_does_not_hide_valid_states(self):
        history = StateHistory(self.work)
        valid = history.new_state_path()
        with contextlib.redirect_stdout(io.StringIO()):
            app.create_snapshot(self.root, valid)
        history.new_state_path().write_bytes(b"broken")
        with zipfile.ZipFile(history.new_state_path(), "w") as archive:
            archive.writestr("metadata.json", "[]")
        items, errors = history.list_states()
        self.assertEqual(len(items), 1)
        self.assertEqual(len(errors), 2)

    def test_worker_queue_is_bounded_on_larger_tree(self):
        for index in range(193):
            (self.root / f"{index}.txt").write_bytes(str(index).encode())
        real_wait = app.wait
        observed = []
        def checked_wait(pending, **kwargs):
            observed.append(len(pending))
            self.assertLessEqual(len(pending), 64)
            return real_wait(pending, **kwargs)
        with patch.object(app, "wait", checked_wait), contextlib.redirect_stdout(io.StringIO()):
            state = app.create_snapshot(self.root, self.work / "bounded.fcrstate", 32)
        self.assertEqual(len(state["files"]), 193)
        self.assertEqual(max(observed), 64)

    def test_legacy_json_and_relative_traversal_rejection(self):
        (self.root / "a.txt").write_bytes(b"a")
        state = self.snapshot("A.fcrstate")
        legacy = self.work / "legacy.json"
        legacy.write_text(json.dumps({**state, "version": 1}), encoding="utf-8")
        self.assertEqual(app.load_snapshot(legacy)["files"], state["files"])
        bad = self.work / "bad.fcrstate"
        write_container(bad, state)
        with zipfile.ZipFile(bad, "a") as archive:
            # Replace the files member in a new archive below, avoiding duplicates.
            metadata = archive.read("metadata.json")
        with zipfile.ZipFile(bad, "w") as archive:
            archive.writestr("metadata.json", metadata)
            archive.writestr("files.json", json.dumps([["../escape", 1, "a" * 64]]))
            archive.writestr("errors.json", "[]")
        with self.assertRaisesRegex(ValueError, "relative path"):
            app.load_snapshot(bad)

    def test_comparison_never_resolves_original_project(self):
        (self.root / "a.txt").write_bytes(b"a")
        state = self.snapshot("A.fcrstate")
        state["root"] = r"\\unavailable-server\project"
        state["files"][0]["path"] = state["root"] + r"\a.txt"
        write_container(self.work / "offline.fcrstate", state)
        with patch.object(app, "inside", side_effect=AssertionError("Live project lookup")):
            with contextlib.redirect_stdout(io.StringIO()):
                counts = app.compare_files(self.work / "offline.fcrstate", self.work / "offline.fcrstate", self.work / "offline.xlsx")
        self.assertEqual(counts, {"Unchanged": 1})

    def test_offline_policy_and_read_only_source(self):
        path = self.root / "a.txt"
        path.write_bytes(b"original")
        with safe_operation((self.root,), self.work):
            self.assertEqual(path.read_bytes(), b"original")
            with self.assertRaises(PermissionError):
                path.write_bytes(b"unexpected")
            with self.assertRaises(PermissionError):
                path.unlink()
            with self.assertRaises(PermissionError):
                socket.socket()
            with self.assertRaises(PermissionError):
                subprocess.Popen(["unused-command"])
            self.assertTrue(app.inside(tempfile.gettempdir(), self.work))
        self.assertEqual(path.read_bytes(), b"original")
        self.assertEqual(list(self.work.glob(".folder-change-review-work-*")), [])

    def test_row_limit_accounts_for_six_header_rows(self):
        self.assertEqual(app.MAX_DATA_ROWS, 1_048_570)


if __name__ == "__main__":
    unittest.main()
