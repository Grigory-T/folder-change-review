"""Size arithmetic, portable placement and exact output-write exceptions."""
import contextlib
import io
import os
import subprocess
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
import folder_change_review as app
import runtime_safety as safety
from state_store import StateHistory

class ReportWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.base=Path(self.temp.name);self.root=self.base/'Project';self.root.mkdir()
        environment=patch.dict(os.environ,{'LOCALAPPDATA':str(self.base/'MachineData')})
        environment.start();self.addCleanup(environment.stop)
    def capture(self,name):
        with contextlib.redirect_stdout(io.StringIO()):return app.create_snapshot(self.root,self.base/(name+'.fcrstate'),32)
    def test_size_arithmetic_and_formats_cover_zero_negative_and_equal_length_edits(self):
        for name,data in [('grow.pdf',b'a'),('delete.pdf',b'123'),('same.docx',b'ab'),('stable',b'')]:
            (self.root/name).write_bytes(data)
        a=self.capture('A')
        (self.root/'grow.pdf').write_bytes(b'a'*1048577);(self.root/'delete.pdf').unlink()
        (self.root/'same.docx').write_bytes(b'CD');(self.root/'add.pdf').write_bytes(b'12345')
        b=self.capture('B');rows=app.compare_snapshots(a,b)
        expected={'grow.pdf':1048576,'delete.pdf':-3,'same.docx':0,'stable':0,'add.pdf':5}
        self.assertEqual({row['FileName']:row['SizeDelta'] for row in rows},expected)
        report=self.root/app.OUTPUT_FILE
        with contextlib.redirect_stdout(io.StringIO()):app.compare_files(self.base/'A.fcrstate',self.base/'B.fcrstate',report)
        book=load_workbook(report);self.addCleanup(book.close);sheet=book['Changes']
        headers=[cell.value for cell in sheet[6]][1:]
        for row in sheet.iter_rows(min_row=7,min_col=2):
            data=dict(zip(headers,row));name=data['File name'].value
            self.assertEqual(data['Size delta (bytes)'].value,expected[name])
            self.assertAlmostEqual(data['Size delta (MiB)'].value,expected[name]/1048576)
            for label in ('A','B','delta'):
                self.assertEqual(data[f'Size {label} (bytes)'].number_format,'#,##0')
                self.assertEqual(data[f'Size {label} (MiB)'].number_format,'#,##0.000')
        self.assertEqual(sheet.freeze_panes,'A7');self.assertEqual(sheet.sheet_view.zoomScale,70)
        # XlsxWriter stores the desired width plus Excel's 5/7 character padding.
        self.assertTrue(all(abs(item.width-5.7109375)<.01 for item in sheet.column_dimensions.values()))
        c=self.capture('C');self.assertEqual(b['files'],c['files'])
        self.assertFalse(list(self.root.glob('.folder-change-review-*')))
    def test_unknown_delta_is_blank_even_with_known_size_on_one_side(self):
        (self.root/'file.pdf').write_bytes(b'file');a=self.capture('A')
        with patch.object(app,'hash_file',side_effect=PermissionError(13,'Denied')),contextlib.redirect_stdout(io.StringIO()):b=app.create_snapshot(self.root,self.base/'B.fcrstate',16)
        row=app.compare_snapshots(a,b)[0]
        self.assertEqual(row['Status'],'Unknown');self.assertIsNone(row['SizeDelta'])
        self.assertEqual(row['SizeA'],4);self.assertIsNone(row['SizeB'])
    def test_artifacts_are_excluded_from_new_and_legacy_comparisons(self):
        (self.root/'doc.pdf').write_bytes(b'doc')
        for name in (app.OUTPUT_FILE,'folder-change-review.exe','old.fcrstate','folder-changes-20261017-000001-abcdef.xlsx'):(self.root/name).write_bytes(b'tool artifact')
        states=self.root/'.states';states.mkdir();(states/'settings.json').write_bytes(b'{}')
        state=self.capture('A');self.assertEqual([r['parts'] for r in state['files']],[['doc.pdf']])
        old={'files':[{'path':str(self.root/app.OUTPUT_FILE),'parts':[app.OUTPUT_FILE],'size':1,'sha256':'a'*64}]}
        self.assertEqual(app.compare_snapshots(old,state)[0]['Status'],'Added')
        self.assertEqual(len(app.compare_snapshots(old,state)),1)
    def test_scoped_output_does_not_allow_other_project_or_network_writes(self):
        document=self.root/'data.pdf';document.write_bytes(b'original');report=self.root/app.OUTPUT_FILE
        with safety.safe_operation((self.root,),self.root):
            self.assertFalse(safety.inside(tempfile.gettempdir(),self.root))
            with app.staged_output(report,local_only=False) as staging:
                Path(staging).write_bytes(b'completed report')
                with self.assertRaises(PermissionError):document.write_bytes(b'bad')
                with self.assertRaises(PermissionError):document.unlink()
                with self.assertRaises(PermissionError):self.root.joinpath('other.xlsx').write_bytes(b'bad')
        self.assertEqual(document.read_bytes(),b'original');self.assertEqual(report.read_bytes(),b'completed report')
        self.assertFalse(safety._output_paths)
    def test_in_project_report_failure_preserves_previous_file_and_cleans_staging(self):
        report=self.root/app.OUTPUT_FILE;report.write_bytes(b'previous')
        with self.assertRaises(OSError),safety.safe_operation((self.root,),self.root):
            with app.staged_output(report,local_only=False) as staging:
                Path(staging).write_bytes(b'incomplete');raise OSError(28,'Disk full')
        self.assertEqual(report.read_bytes(),b'previous')
        self.assertFalse(list(self.root.glob('.folder-change-review-*')))
    def test_fixed_default_cli_report_and_quoted_root(self):
        (self.root/'doc.pdf').write_bytes(b'doc');self.capture('A');self.capture('B')
        with patch.object(app,'application_directory',return_value=self.root),contextlib.redirect_stdout(io.StringIO()):
            code=app.main(['compare','--before',str(self.base/'A.fcrstate'),'--after',str(self.base/'B.fcrstate')])
        self.assertEqual(code,0);self.assertTrue((self.root/app.OUTPUT_FILE).exists())
        with contextlib.redirect_stdout(io.StringIO()):app.create_snapshot("'"+str(self.root)+"'",self.base/'quoted.fcrstate',1)
        for raw,expected in [(r'"\\server\share\a b"',r'\\server\share\a b'),(" 'C:\\A B' ",'C:\\A B'),('"unmatched','"unmatched')]:
            self.assertEqual(safety.pasted_path(raw),expected)
    def test_alias_output_and_relative_issue_paths_preserve_literal_names(self):
        alias=self.base/'alias'
        if os.name=='nt':
            subprocess.run(['cmd','/c','mklink','/J',str(alias),str(self.root)],check=True,capture_output=True)
        else:
            alias.symlink_to(self.root,target_is_directory=True)
        with safety.safe_operation((self.root,),alias),app.staged_output(alias/app.OUTPUT_FILE,local_only=False) as temp:Path(temp).write_bytes(b'report')
        self.assertEqual((self.root/app.OUTPUT_FILE).read_bytes(),b'report')
        target=self.root/'business.xlsx';target.write_bytes(b'keep');link=self.base/'link.xlsx'
        if os.name=='nt':
            with patch.object(app.os.path,'islink',return_value=True), self.assertRaises(ValueError):
                with app.staged_output(link,local_only=False):pass
        else:
            link.symlink_to(target)
            with self.assertRaises(ValueError):
                with app.staged_output(link,local_only=False):pass
        self.assertEqual(target.read_bytes(),b'keep')
    def test_network_application_uses_local_machine_history(self):
        original=safety.local_storage
        with patch('state_store.local_storage',return_value=False):history=StateHistory(self.base/'NetworkTool')
        self.assertNotEqual(history.states.parent,history.directory)
        self.assertTrue(original(history.states))
        self.assertEqual(history.new_report_path().name,app.OUTPUT_FILE)
    def test_operation_guard_and_output_error_explanation(self):
        with safety.safe_operation((),self.base):
            with self.assertRaises(RuntimeError):
                with safety.safe_operation((),self.base):pass
        self.assertIn('Close it in Excel',app.report_error_message(PermissionError(13,'locked')))
        self.assertFalse(safety._output_paths)

    def test_local_tool_history_is_outside_project_and_legacy_import_is_read_only(self):
        tool=self.root/'Tool';tool.mkdir();legacy=tool/'.states';legacy.mkdir()
        (self.root/'document.pdf').write_bytes(b'original')
        state=self.capture('A')
        original=(self.base/'A.fcrstate').read_bytes()
        old=legacy/'legacy.fcrstate';old.write_bytes(original)
        settings=legacy/'settings.json';settings.write_text('{"workers":32}')
        history=StateHistory(tool)
        self.assertFalse(safety.inside(history.states,self.root))
        self.assertEqual(history.settings()['workers'],32)
        items,errors=history.list_states()
        self.assertEqual([item['path'] for item in items],[old]);self.assertEqual(errors,[])
        self.assertFalse(history.states.exists())  # startup reads metadata only
        history.save_settings({'workers':16})
        self.assertEqual(old.read_bytes(),original)
        self.assertEqual(settings.read_text(),'{"workers":32}')
        self.assertEqual((history.states/old.name).read_bytes(),original)
        self.assertEqual(history.settings()['workers'],16)
        self.assertEqual(len(history.list_states()[0]),1)
        with contextlib.redirect_stdout(io.StringIO()):app.create_snapshot(self.root,history.new_state_path())
        self.assertEqual(app.load_snapshot(history.list_states()[0][0]['path'])['files'],state['files'])

    def test_legacy_copy_failure_keeps_original_history_usable(self):
        tool=self.root/'Tool';(tool/'.states').mkdir(parents=True)
        self.capture('A');old=tool/'.states'/'legacy.fcrstate'
        old.write_bytes((self.base/'A.fcrstate').read_bytes())
        history=StateHistory(tool)
        with patch('state_store.shutil.copyfile',side_effect=PermissionError('Locked')):
            history.save_settings({'workers':16})
        self.assertEqual(history.list_states()[0][0]['path'],old)
        self.assertFalse(list(history.states.glob('.folder-change-review-*')))
