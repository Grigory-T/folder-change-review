"""Truthful comparison under file locks, permissions and incomplete enumeration."""
import copy
import contextlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from openpyxl import load_workbook
import folder_change_review as app
from state_store import read_metadata


class PartialStateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.root = self.base/'Project'; self.root.mkdir()
        for directory in ('Area', 'Area2'):
            (self.root/directory).mkdir()
            (self.root/directory/'report.pdf').write_bytes(b'document')
        with contextlib.redirect_stdout(io.StringIO()):
            self.complete = app.create_snapshot(self.root, self.base/'A.fcrstate', 1)

    def partial(self, kind='directory', area='Area'):
        state = copy.deepcopy(self.complete)
        target = self.root/area
        if kind == 'file': target /= 'report.pdf'
        state['files'] = [record for record in state['files'] if record['parts'][0] != area]
        state['file_count'] = len(state['files'])
        state.update(complete=False, read_error_count=1, read_errors=[app.read_issue(str(target), list(target.relative_to(self.root).parts), kind, PermissionError(13, 'Access denied'))])
        return state

    def test_coverage_matrix_preserves_known_changes_and_unknown_absence(self):
        for kind in ('directory', 'unknown', 'file'):
            for a_partial, b_partial in ((False,True),(True,False),(True,True)):
                with self.subTest(kind=kind,A=a_partial,B=b_partial):
                    a = self.partial(kind) if a_partial else self.complete
                    b = self.partial(kind) if b_partial else self.complete
                    rows = app.compare_snapshots(a,b)
                    statuses = {row['FullPath']:row['Status'] for row in rows}
                    self.assertEqual(statuses[str(self.root/'Area2/report.pdf')], 'Unchanged')
                    if kind == 'file' or not (a_partial and b_partial):
                        self.assertEqual(statuses[str(self.root/'Area/report.pdf')], 'Unknown')
                    else:
                        # Neither state knows the names inside a denied folder.
                        self.assertNotIn(str(self.root/'Area/report.pdf'), statuses)

    def test_prefix_boundaries_do_not_hide_real_additions_or_deletions(self):
        partial = self.partial()
        partial['files'] = []
        rows = app.compare_snapshots(self.complete,partial)
        self.assertEqual({row['FileName']:row['Status'] for row in rows if 'Area2' in row['parts']}, {'report.pdf':'Deleted'})
        rows = app.compare_snapshots(partial,self.complete)
        self.assertEqual([row['Status'] for row in rows if 'Area2' in row['parts']], ['Added'])

    def test_readable_records_take_precedence_over_interrupted_folder_scope(self):
        partial = self.partial()
        partial['files'] = copy.deepcopy(self.complete['files'])
        partial['files'][0]['sha256'] = '0'*64
        rows = app.compare_snapshots(self.complete,partial)
        self.assertEqual(sorted(row['Status'] for row in rows), ['Changed','Unchanged'])

    def test_file_failure_without_previous_inventory_has_an_unknown_leaf(self):
        empty = copy.deepcopy(self.complete); empty['files'] = []
        state = self.partial('file')
        rows = app.compare_snapshots(empty,state)
        row = next(row for row in rows if row['Status']=='Unknown')
        self.assertIsNone(row['SizeA']); self.assertIsNone(row['SizeB'])
        self.assertEqual(row['FileName'],'report.pdf')

    def test_partial_report_has_warning_and_literal_issue_details(self):
        partial = self.partial('file')
        partial['read_errors'][0]['message'] = '=1+1'
        output = self.base/'partial.fcrstate'
        with app.staged_output(output) as staged: app.write_container(staged, partial)
        with contextlib.redirect_stdout(io.StringIO()):
            counts = app.compare_files(self.base/'A.fcrstate',output,self.base/'Changes.xlsx')
        self.assertEqual(counts, {'Unknown':1,'Unchanged':1})
        book = load_workbook(self.base/'Changes.xlsx'); self.addCleanup(book.close)
        self.assertEqual(book.sheetnames,['Changes','Read issues'])
        self.assertIn('Incomplete coverage',book['Changes']['B3'].value)
        self.assertEqual(book['Read issues']['E2'].value,'=1+1')
        self.assertEqual(book['Read issues']['E2'].data_type,'s')
        headers = [cell.value for cell in book['Changes'][6]][1:]
        unknown = next(dict(zip(headers,row)) for row in book['Changes'].iter_rows(min_row=7,min_col=2,values_only=True) if 'Unknown' in row)
        self.assertEqual([unknown[key] for key in ('Added','Deleted','Changed')],[None]*3)
        self.assertEqual(unknown['last'],1)
        self.assertIsNotNone(unknown['Size A (MiB)'])
        self.assertIsNone(unknown['Size B (MiB)'])

    def test_root_read_failure_remains_fatal_and_preserves_output(self):
        output = self.base/'A.fcrstate'; original=output.read_bytes()
        original_scandir = app.os.scandir
        def denied(path):
            if not isinstance(path, int) and app.absolute(path) == str(self.root):
                raise PermissionError(13,'Denied',str(self.root))
            return original_scandir(path)
        with patch.object(app.os,'scandir',side_effect=denied):
            with self.assertRaises(OSError): app.create_snapshot(self.root,output,1)
        self.assertEqual(output.read_bytes(), original)

    def test_old_formats_still_load_but_partial_format_cannot_be_downgraded(self):
        state=copy.deepcopy(self.complete)
        for version in (1,2):
            state['version']=version
            state.pop('read_errors',None);state.pop('read_error_count',None);state.pop('complete',None)
            path=self.base/f'legacy{version}.json';path.write_text(json.dumps(state))
            self.assertEqual(app.load_snapshot(path)['files'],self.complete['files'])
        partial=self.partial();partial['version']=2
        path=self.base/'downgraded.json';path.write_text(json.dumps(partial))
        with self.assertRaises(ValueError):app.load_snapshot(path)

    def test_corrupt_coverage_metadata_or_paths_are_rejected(self):
        for change in ('count','complete','outside','kind','missing-errors'):
            with self.subTest(change=change):
                state=self.partial()
                if change=='count':state['read_error_count']=0
                elif change=='complete':state['complete']=True
                elif change=='outside':state['read_errors'][0]['parts']=['..','other']
                elif change=='kind':state['read_errors'][0]['kind']='invalid'
                else:state.pop('read_errors')
                path=self.base/f'{change}.json';path.write_text(json.dumps(state))
                with self.assertRaises(ValueError):app.load_snapshot(path)

    def test_history_metadata_does_not_load_the_error_inventory(self):
        state=self.partial();path=self.base/'partial.fcrstate'
        with app.staged_output(path) as staged:app.write_container(staged,state)
        original=zipfile.ZipFile.read
        def metadata_only(archive,name,*args):
            self.assertEqual(name.filename if isinstance(name,zipfile.ZipInfo) else name,'metadata.json')
            return original(archive,name,*args)
        with patch.object(zipfile.ZipFile,'read',metadata_only):
            metadata=read_metadata(path)
        self.assertFalse(metadata['complete']);self.assertEqual(metadata['read_error_count'],1)

    def test_windows_unc_scopes_match_ancestors_without_matching_siblings(self):
        root = r'\\server\share\Project'
        before = {'root':root,'files':[
            {'path':root+r'\Area\Nested\one.pdf','parts':['Area','Nested','one.pdf'],'size':1,'sha256':'a'*64},
            {'path':root+r'\Area2\two.pdf','parts':['Area2','two.pdf'],'size':1,'sha256':'a'*64}]}
        after = {'root':root,'files':[],'read_errors':[{'path':root+r'\Area','parts':['Area'],'kind':'directory'}]}
        rows=app.compare_snapshots(before,after)
        self.assertEqual({row['FileName']:row['Status'] for row in rows},{'one.pdf':'Unknown','two.pdf':'Deleted'})

    def test_disk_full_during_save_preserves_output_and_cleans_staging(self):
        output=self.base/'A.fcrstate';original=output.read_bytes()
        def disk_full(path,*args):
            Path(path).write_bytes(b'incomplete output')
            raise OSError(28,'No space left on device')
        with patch.object(app,'write_container',side_effect=disk_full):
            with self.assertRaises(OSError):app.create_snapshot(self.root,output,32)
        self.assertEqual(output.read_bytes(),original)
        self.assertEqual(list(self.base.glob('.folder-change-review-*')),[])

    def test_two_unreadable_states_report_coverage_even_without_known_filenames(self):
        state=self.partial();state['files']=[];state['file_count']=0
        path=self.base/'partial.fcrstate'
        with app.staged_output(path) as staged:app.write_container(staged,state)
        with contextlib.redirect_stdout(io.StringIO()):app.compare_files(path,path,self.base/'Empty-coverage.xlsx')
        book=load_workbook(self.base/'Empty-coverage.xlsx');self.addCleanup(book.close)
        self.assertEqual(book['Changes'].max_row,6)
        self.assertIn('Incomplete coverage',book['Changes']['B3'].value)
        self.assertEqual(book['Read issues'].max_row,3)
